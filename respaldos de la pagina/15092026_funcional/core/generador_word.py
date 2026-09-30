# -*- coding: utf-8 -*-
"""
Created on Fri Jul 31 17:03:16 2026

@author: dchable
"""

# core/generador_word.py
import pandas as pd
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from core.word_utils import (dar_formato_celda, forzar_bordes_tabla, borrar_borde_celda, 
                             buscar_parrafos_tabla, limpiar_tablas_viejas)

def inyectar_tabla_vertices_en_word(ruta_doc, clave_ac, df_vertices):
    try:
        if df_vertices is None or df_vertices.empty: return "⚠️ Vertices.xlsx no cargado."
        
        df_verts = df_vertices[df_vertices['ID_ACUIFERO'] == clave_ac].copy()
        if df_verts.empty: return "⚠️ El Acuífero no tiene vértices registrados."
        
        # Lógica de cierre de polígono
        df_verts = df_verts.sort_values(by='VERTICE')
        fila_cierre = df_verts[df_verts.duplicated(subset=['VERTICE'], keep='first')]
        df_verts = df_verts.drop_duplicates(subset=['VERTICE'], keep='first')
        df_verts = pd.concat([df_verts, fila_cierre])
        df_verts['OBSERVACIONES'] = df_verts['OBSERVACIONES'].fillna('')
        
        doc = Document(ruta_doc)
        parrafos = buscar_parrafos_tabla(doc, ["tabla", "coordenadas"])
        if not parrafos: return "⚠️ No se encontró título Tabla 1 (Vértices)."
        
        table = doc.add_table(rows=2, cols=8)
        forzar_bordes_tabla(table)
        
        row0 = table.rows[0].cells
        row1 = table.rows[1].cells
        
        dar_formato_celda(row0[0].merge(row1[0]), "VÉRTICE", "C0D7EE", True, 9, color_texto="244062")
        dar_formato_celda(row0[1].merge(row0[3]), "LONGITUD OESTE", "C0D7EE", True, 9, color_texto="244062")
        dar_formato_celda(row0[4].merge(row0[6]), "LATITUD NORTE", "C0D7EE", True, 9, color_texto="244062")
        dar_formato_celda(row0[7].merge(row1[7]), "OBSERVACIONES", "C0D7EE", True, 9, color_texto="244062")
        
        for i, val in enumerate(["GRADOS", "MINUTOS", "SEGUNDOS", "GRADOS", "MINUTOS", "SEGUNDOS"]):
            dar_formato_celda(row1[i+1], val, "C0D7EE", True, 8, color_texto="244062")
            
        def a_float(val):
            try: return abs(float(val))
            except: return 0.0

        for _, row in df_verts.iterrows():
            nueva_fila = table.add_row()
            nueva_fila.height = Cm(0.43)
            cells = nueva_fila.cells
            dar_formato_celda(cells[0], str(row.get('VERTICE', '')), None, False, 9)
            dar_formato_celda(cells[1], str(row.get('LONG_G', '')), None, False, 9)
            dar_formato_celda(cells[2], str(row.get('LONG_M', '')), None, False, 9)
            dar_formato_celda(cells[3], f"{a_float(row.get('LONG_S', 0)):.1f}".replace("-0.0", "0.0"), None, False, 9)
            dar_formato_celda(cells[4], str(row.get('LAT_G', '')), None, False, 9)
            dar_formato_celda(cells[5], str(row.get('LAT_M', '')), None, False, 9)
            dar_formato_celda(cells[6], f"{a_float(row.get('LAT_S', 0)):.1f}".replace("-0.0", "0.0"), None, False, 9)
            dar_formato_celda(cells[7], str(row.get('OBSERVACIONES', '')), None, False, 8)

        parrafos[-1]._p.addnext(table._tbl)
        limpiar_tablas_viejas(parrafos)
        doc.save(ruta_doc)
        return "✅ Vértices inyectados."
    except Exception as e: return f"❌ Error Vértices: {e}"

def inyectar_tabla_flujo_en_word(ruta_doc, clave_ac, df_global, palabras_clave, nombre_log):
    try:
        if df_global is None: return f"⚠️ Parquet no encontrado para {nombre_log}."
        df_flujo = df_global[df_global['CLV_ACUI'] == clave_ac].copy()
        if df_flujo.empty: return f"⚠️ Sin registros para {nombre_log}."

        doc = Document(ruta_doc)
        parrafos = buscar_parrafos_tabla(doc, palabras_clave)
        if not parrafos: return f"⚠️ Título de {nombre_log} no encontrado."

        table = doc.add_table(rows=1, cols=8)
        table.autofit = True
        forzar_bordes_tabla(table)
        
        encabezados = ["CELDA", "LONGITUD B\n(m)", "ANCHO a\n(m)", "h2-h1\n(m)", "Gradiente i", "T\n(m²/s)", "CAUDAL Q\n(m³/s)", "VOLUMEN\n(hm³/año)"]
        for i, texto in enumerate(encabezados):
            dar_formato_celda(table.rows[0].cells[i], texto, fondo_color="C0D7EE", negrita=True, tamano=10, color_texto="244062")
        
        def fmt_num(val, dec):
            if pd.isna(val) or str(val).strip() in ["", "nan", "None"]: return "-"
            try: return f"{float(val):.{dec}f}" if dec > 0 else f"{int(round(float(val)))}"
            except: return str(val)

        total_vol = 0.0
        for _, row in df_flujo.iterrows():
            nueva_fila = table.add_row()
            cells = nueva_fila.cells
            
            # Filtrar columnas no deseadas
            datos_validos = [v for k, v in row.items() if not any(x in str(k).lower() for x in ['clv_acui', 'unnamed', 'index'])]
            while len(datos_validos) < 8: datos_validos.append("-")
            
            val_vol = datos_validos[7]
            try:
                if pd.notna(val_vol) and str(val_vol).lower() not in ['nan', 'none', '-']: total_vol += float(val_vol)
            except: pass
            
            dar_formato_celda(cells[0], str(datos_validos[0]).replace('.0', '') if str(datos_validos[0]).lower() not in ['nan','none',''] else "-")
            dar_formato_celda(cells[1], fmt_num(datos_validos[1], 0))
            dar_formato_celda(cells[2], fmt_num(datos_validos[2], 0))
            dar_formato_celda(cells[3], fmt_num(datos_validos[3], 0))
            dar_formato_celda(cells[4], fmt_num(datos_validos[4], 5))
            dar_formato_celda(cells[5], fmt_num(datos_validos[5], 4))
            dar_formato_celda(cells[6], fmt_num(datos_validos[6], 4))
            dar_formato_celda(cells[7], fmt_num(val_vol, 1))

        fila_total = table.add_row()
        celda_label = fila_total.cells[0].merge(fila_total.cells[6])
        dar_formato_celda(celda_label, "TOTAL", fondo_color="C0D7EE", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.RIGHT, color_texto="244062")
        dar_formato_celda(fila_total.cells[7], fmt_num(total_vol, 1), fondo_color="C0D7EE", negrita=True, color_texto="244062")
        borrar_borde_celda(celda_label, 'right')
        borrar_borde_celda(fila_total.cells[7], 'left')

        parrafos[-1]._p.addnext(table._tbl)
        limpiar_tablas_viejas(parrafos)
        doc.save(ruta_doc)
        return f"✅ {nombre_log} inyectadas."
    except Exception as e: return f"❌ Error {nombre_log}: {e}"

def inyectar_tabla_almacenamiento_en_word(ruta_doc, clave_ac, df_almacenamiento):
    try:
        if df_almacenamiento is None: return "⚠️ Parquet de Almacenamiento no encontrado."
        df_datos = df_almacenamiento[df_almacenamiento['CLV_ACUI'] == clave_ac].copy()
        if df_datos.empty: return f"⚠️ Sin registros de almacenamiento."

        doc = Document(ruta_doc)
        parrafos = buscar_parrafos_tabla(doc, ["tabla", "cambio de almacenamiento"])
        if not parrafos: return "⚠️ Título Almacenamiento no encontrado."

        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        forzar_bordes_tabla(table)
        
        encabezados = ["Evolución\n(m)", "Evolución media\n(m)", "Área\n(km²)", "Sy", "ΔV(S)\n(hm³/año)"]
        for i, texto in enumerate(encabezados):
            dar_formato_celda(table.rows[0].cells[i], texto, fondo_color="C0D7EE", negrita=True, tamano=10, color_texto="244062")

        total_vol = 0.0
        val_promedio_anual = None
        
        for _, row in df_datos.iterrows():
            evo_str = str(row.get('EVOLUCION_m', ''))
            if "promedio anual" in evo_str.lower():
                val_promedio_anual = row.get('VOLUMEN_hm3')
                continue 
            
            nueva_fila = table.add_row()
            cells = nueva_fila.cells
            
            def fmt(val, dec):
                if pd.isna(val) or str(val).lower() in ['nan', 'none', '', '-']: return "-"
                try: return f"{float(val):.{dec}f}"
                except: return str(val)

            dar_formato_celda(cells[0], fmt(row.get('EVOLUCION_m'), 2))
            dar_formato_celda(cells[1], fmt(row.get('EVOLUCION_MEDIA_m'), 2))
            dar_formato_celda(cells[2], fmt(row.get('AREA_km2'), 2))
            dar_formato_celda(cells[3], fmt(row.get('Sy'), 3))
            vol = row.get('VOLUMEN_hm3')
            dar_formato_celda(cells[4], fmt(vol, 1))
            try: total_vol += float(vol)
            except: pass

        # Fila Total
        fila_total = table.add_row()
        celda_label = fila_total.cells[0].merge(fila_total.cells[3])
        dar_formato_celda(celda_label, "TOTAL", fondo_color="C0D7EE", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.RIGHT, color_texto="244062")
        dar_formato_celda(fila_total.cells[4], fmt(total_vol, 1), fondo_color="C0D7EE", negrita=True, color_texto="244062")
        borrar_borde_celda(celda_label, 'right')
        borrar_borde_celda(fila_total.cells[4], 'left')

        # Fila Promedio Anual (si existe)
        if val_promedio_anual is not None and str(val_promedio_anual).lower() not in ['nan', 'none', '', '-']:
            fila_prom = table.add_row()
            celda_label_prom = fila_prom.cells[0].merge(fila_prom.cells[3])
            dar_formato_celda(celda_label_prom, "Promedio anual", fondo_color="C0D7EE", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.RIGHT, color_texto="244062")
            dar_formato_celda(fila_prom.cells[4], fmt(val_promedio_anual, 1), fondo_color="C0D7EE", negrita=True, color_texto="244062")
            borrar_borde_celda(celda_label_prom, 'right')
            borrar_borde_celda(fila_prom.cells[4], 'left')

        parrafos[-1]._p.addnext(table._tbl)
        limpiar_tablas_viejas(parrafos)
        doc.save(ruta_doc)
        return "✅ Almacenamiento inyectado."
    except Exception as e: return f"❌ Error Almacenamiento: {e}"

def inyectar_tabla_evapotranspiracion_en_word(ruta_doc, clave_ac, df_etr):
    try:
        if df_etr is None: return "⚠️ Parquet Evapotranspiración no encontrado."
        df_datos = df_etr[df_etr['CLV_ACUI'] == clave_ac].copy()
        if df_datos.empty: return f"⚠️ Sin registros de evapotranspiración."

        doc = Document(ruta_doc)
        parrafos = buscar_parrafos_tabla(doc, ["tabla", "evapotranspiracion"])
        if not parrafos: return "⚠️ Título ETR no encontrado."

        table = doc.add_table(rows=1, cols=7)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        forzar_bordes_tabla(table)
        
        encabezados = ["RANGOS DE\nPROFUNDIDAD\n(m)", "PROFUNDIDAD\nMEDIA\n(m)", "ÁREA\n(km²)", "LÁMINA\nETR\n(m)", "PROFUNDIDAD MÁXIMA\nDE EXTINCIÓN DE LA\nETR", "% ETR", "VOLUMEN ETR\n(hm³/año)"]
        for i, texto in enumerate(encabezados):
            dar_formato_celda(table.rows[0].cells[i], texto, fondo_color="C0D7EE", negrita=True, tamano=9, color_texto="244062")
        
        total_area, total_vol = 0.0, 0.0
        columnas_db = ["RANGOS_DE_PROFUNDIDAD_m", "PROFUNDIDAD_MEDIA_m", "AREA_km2", "LAMINA_ETR_m", "PROF_MAX_EXTINCION_ETR", "PORCENTAJE_ETR", "VOLUMEN_ETR_hm3_ano"]
        
        for _, row in df_datos.iterrows():
            nueva_fila = table.add_row()
            cells = nueva_fila.cells
            
            def fmt(val, col_idx):
                val_str = str(val).strip()
                if val_str.lower() in ['nan', 'none', '', '-']: return "-"
                try:
                    f_val = float(val_str)
                    formatos = {2: "{:.1f}", 3: "{:.4f}", 4: "{:.0f}", 5: "{:.1f}", 6: "{:.1f}"}
                    return formatos.get(col_idx, "{:.1f}").format(f_val)
                except: return val_str

            for i, col_name in enumerate(columnas_db):
                dar_formato_celda(cells[i], fmt(row.get(col_name), i))
            
            try: total_area += float(str(row.get('AREA_km2')).strip())
            except: pass
            try: total_vol += float(str(row.get('VOLUMEN_ETR_hm3_ano')).strip())
            except: pass

        fila_total = table.add_row()
        celda_total_1 = fila_total.cells[0].merge(fila_total.cells[1])
        celda_area = fila_total.cells[2]
        celda_total_2 = fila_total.cells[3].merge(fila_total.cells[5])
        celda_vol = fila_total.cells[6]
        
        dar_formato_celda(celda_total_1, "TOTAL", fondo_color="C0D7EE", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.RIGHT, color_texto="244062")
        dar_formato_celda(celda_area, f"{total_area:.1f}" if total_area >= 0 else "-", fondo_color="C0D7EE", negrita=True, color_texto="244062")
        dar_formato_celda(celda_total_2, "TOTAL", fondo_color="C0D7EE", negrita=True, alineacion=WD_ALIGN_PARAGRAPH.RIGHT, color_texto="244062")
        dar_formato_celda(celda_vol, f"{total_vol:.1f}" if total_vol >= 0 else "-", fondo_color="C0D7EE", negrita=True, color_texto="244062")

        bordes_a_quitar = [(celda_total_1, 'right'), (celda_area, 'left'), (celda_area, 'right'), (celda_total_2, 'left'), (celda_total_2, 'right'), (celda_vol, 'left')]
        for celda_obj, lado_borde in bordes_a_quitar:
            borrar_borde_celda(celda_obj, lado_borde)

        parrafos[-1]._p.addnext(table._tbl)
        limpiar_tablas_viejas(parrafos)
        doc.save(ruta_doc)
        return "✅ ETR inyectada."
    except Exception as e: return f"❌ Error ETR: {e}"

def modificar_censo_y_bombeo(ruta_doc, p_aprov, p_vol):
    try:
        doc = Document(ruta_doc)
        body = doc._body._body
        borrando_censo, borrando_bombeo = False, False
        p_bombeo_objetivo, elementos_a_borrar = None, []
        
        for element in body:
            es_encabezado_paro = False
            txt = ""
            if element.tag.endswith('}p'):
                p = Paragraph(element, doc)
                txt = p.text.strip().upper()
                style_lower = p.style.name.lower() if p.style else ""
                
                paro_por_estilo = any(h in style_lower for h in ["heading", "título", "titulo", "subtítulo", "subtitulo", "sub-sub"])
                conceptos_freno = ["EVAPOTRANSPIRACIÓN", "EVAPOTRANSPIRACION", "ETR", "SALIDA", "SALIDAS", "BALANCE DE AGUAS", "DISPONIBILIDAD", "BIBLIOGRAFÍA", "DESCARGA", "CAMBIO DE ALMACENAMIENTO"]
                paro_por_texto = any(c in txt for c in conceptos_freno) and len(txt) < 110
                
                es_encabezado_paro = paro_por_estilo or paro_por_texto
                
                if "CENSO DE APROVECHAMIENTOS" in txt and len(txt) < 100:
                    borrando_censo, borrando_bombeo = True, False
                    elementos_a_borrar.append(element)
                    continue
                    
                if "BOMBEO" in txt and any(h in style_lower for h in ["sub-sub", "heading 3", "heading3"]):
                    p_bombeo_objetivo = p
                    borrando_bombeo, borrando_censo = True, False
                    continue 
            
            if es_encabezado_paro:
                borrando_censo = borrando_bombeo = False
                
            if borrando_censo or borrando_bombeo:
                elementos_a_borrar.append(element)

        for el in elementos_a_borrar:
            if el.getparent() is not None:
                el.getparent().remove(el)
                
        if p_bombeo_objetivo:
            p1 = doc.add_paragraph()
            p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run1 = p1.add_run(p_aprov)
            run1.font.name, run1.font.size = 'Noto Sans', Pt(11)
            
            p_espacio = doc.add_paragraph()
            
            p2 = doc.add_paragraph()
            p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run2 = p2.add_run(p_vol)
            run2.font.name, run2.font.size = 'Noto Sans', Pt(11)
            
            p_bombeo_objetivo._p.addnext(p1._p)
            p1._p.addnext(p_espacio._p)
            p_espacio._p.addnext(p2._p)
        
        doc.save(ruta_doc)
        return "✅ Censo/Bombeo actualizados."
    except Exception as e: return f"❌ Error Censo/Bombeo: {e}"