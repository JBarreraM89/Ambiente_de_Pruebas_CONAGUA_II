# -*- coding: utf-8 -*-
"""
Created on Mon Aug  3 11:56:27 2026

@author: dchable
"""

# core/word_utils.py
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def dar_formato_celda(celda, texto, fondo_color=None, negrita=False, tamano=9, alineacion=WD_ALIGN_PARAGRAPH.CENTER, color_texto=None):
    """Aplica formato profesional (color, negritas, alineación) a una celda de Word."""
    tcPr = celda._tc.get_or_add_tcPr()
    
    # Centrado vertical
    tcVAlign = tcPr.find(qn('w:vAlign'))
    if tcVAlign is None:
        tcVAlign = OxmlElement('w:vAlign')
        tcPr.append(tcVAlign)
    tcVAlign.set(qn('w:val'), "center")
    
    # Autoajuste de ancho
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), '0')
    tcW.set(qn('w:type'), 'auto')
    
    # Fondo de celda
    if fondo_color:
        shd = tcPr.find(qn('w:shd'))
        if shd is not None: tcPr.remove(shd)
        tcShd = OxmlElement('w:shd')
        tcShd.set(qn('w:val'), 'clear')
        tcShd.set(qn('w:color'), 'auto')
        tcShd.set(qn('w:fill'), fondo_color)
        tcPr.append(tcShd)
        
    # Limpiar y escribir texto
    celda.text = "" 
    p = celda.paragraphs[0]
    p.alignment = alineacion
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    run = p.add_run(texto)
    run.font.name = 'Noto Sans'
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    
    if color_texto:
        run.font.color.rgb = RGBColor.from_string(color_texto)

def forzar_bordes_tabla(tbl):
    """Dibuja bordes negros sólidos institucionales en una tabla de docx."""
    tblPr = tbl._tbl.tblPr
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    tblBorders = OxmlElement('w:tblBorders')
    for borde in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement(f'w:{borde}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4') 
        element.set(qn('w:color'), '000000') 
        tblBorders.append(element)
    tblPr.append(tblBorders)

def borrar_borde_celda(celda_obj, lado_borde):
    """Cirugía XML: Borra un borde específico para simular fusión de celdas compleja."""
    tcPr = celda_obj._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders')) or OxmlElement('w:tcBorders')
    borde_xml = OxmlElement(f'w:{lado_borde}')
    borde_xml.set(qn('w:val'), 'nil')
    tcBorders.append(borde_xml)
    tcPr.append(tcBorders)

def limpiar_str_word(texto):
    """Limpia acentos y mayúsculas para buscar títulos de tablas."""
    return texto.lower().replace("á","a").replace("é","e").replace("í","i").replace("ó","o").replace("ú","u")

def buscar_parrafos_tabla(doc, palabras_clave):
    """Busca en el documento párrafos que contengan palabras clave específicas."""
    parrafos_encontrados = []
    for p in doc.paragraphs:
        texto_p = limpiar_str_word(p.text).strip()
        if texto_p.startswith("tabla") and len(texto_p) < 150 and all(limpiar_str_word(pal) in texto_p for pal in palabras_clave):
            parrafos_encontrados.append(p)
    return parrafos_encontrados

def limpiar_tablas_viejas(parrafos_encontrados):
    """Borra párrafos basura dejados por las plantillas anteriores."""
    for p_basura in parrafos_encontrados[:-1]:
        xml_str = p_basura._element.xml
        if 'w:type="page"' in xml_str or 'pageBreakBefore' in xml_str:
            for t in p_basura._element.iter():
                if t.tag.endswith('}t'): t.text = ''
        else:
            if p_basura._element.getparent() is not None:
                p_basura._element.getparent().remove(p_basura._element)