# -*- coding: utf-8 -*-
# core/word_utils.py
import pandas as pd
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def dar_formato_celda(celda, texto, fondo_color=None, negrita=False, tamano=9, alineacion=WD_ALIGN_PARAGRAPH.CENTER, color_texto=None):
    """Aplica formato profesional (color, negritas, alineación) a una celda de Word."""
    tcPr = celda._tc.get_or_add_tcPr()
    
    tcVAlign = tcPr.find(qn('w:vAlign'))
    if tcVAlign is None:
        tcVAlign = OxmlElement('w:vAlign')
        tcPr.append(tcVAlign)
    tcVAlign.set(qn('w:val'), "center")
    
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), '0')
    tcW.set(qn('w:type'), 'auto')
    
    if fondo_color:
        shd = tcPr.find(qn('w:shd'))
        if shd is not None:
            tcPr.remove(shd)
        tcShd = OxmlElement('w:shd')
        tcShd.set(qn('w:val'), 'clear')
        tcShd.set(qn('w:color'), 'auto')
        tcShd.set(qn('w:fill'), fondo_color)
        tcPr.append(tcShd)
        
    celda.text = "" 
    p = celda.paragraphs[0]
    p.alignment = alineacion
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    run = p.add_run(texto)
    run.font.name = 'Noto Sans'
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    
    if color_texto:
        run.font.color.rgb = RGBColor.from_string(color_texto)

def configurar_encabezado_repetible_y_paginacion(table, num_filas_encabezado=1):
    """
    1. Hace que las primeras `num_filas_encabezado` se repitan automáticamente si la tabla
       se divide en la siguiente página (<w:tblHeader/>).
    2. Evita que cualquier fila se corte a la mitad entre dos páginas (<w:cantSplit/>).
    """
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        
        # Evitar que una fila se parta a la mitad en un salto de página
        if trPr.find(qn('w:cantSplit')) is None:
            trPr.append(OxmlElement('w:cantSplit'))
            
        # Repetir fila(s) de título en cada página nueva
        if r_idx < num_filas_encabezado:
            if trPr.find(qn('w:tblHeader')) is None:
                tblHeader = OxmlElement('w:tblHeader')
                tblHeader.set(qn('w:val'), 'true')
                trPr.append(tblHeader)

def forzar_bordes_tabla(tbl):
    """Dibuja bordes negros sólidos institucionales y activa la repetición de encabezados."""
    tblPr = tbl._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is not None:
        tblPr.remove(tblW)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '5000')
    tblW.set(qn('w:type'), 'pct')
    tblPr.append(tblW)

    tblBorders = tblPr.find(qn('w:tblBorders'))
    if tblBorders is not None:
        tblPr.remove(tblBorders)
    tblBorders = OxmlElement('w:tblBorders')
    for borde in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        element = OxmlElement(f'w:{borde}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4') 
        element.set(qn('w:color'), '000000') 
        tblBorders.append(element)
    tblPr.append(tblBorders)
    num_headers = 2 if len(tbl.rows) == 2 else 1
    configurar_encabezado_repetible_y_paginacion(tbl, num_filas_encabezado=num_headers)

def borrar_borde_celda(celda_obj, lado_borde):
    """Cirugía XML: Borra un borde específico para simular fusión de celdas compleja."""
    tcPr = celda_obj._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    borde_xml = OxmlElement(f'w:{lado_borde}')
    borde_xml.set(qn('w:val'), 'nil')
    tcBorders.append(borde_xml)

def limpiar_str_word(texto):
    """Limpia acentos y mayúsculas para buscar títulos de tablas."""
    return texto.lower().replace("á","a").replace("é","e").replace("í","i").replace("ó","o").replace("ú","u")

def buscar_parrafos_tabla(doc, palabras_clave):
    """
    Busca en el documento párrafos de título de tabla.
    Soporta palabras simples o tuplas de sinónimos: ej. ["tabla", "entradas", ("flujo", "subterraneas")]
    """
    parrafos_encontrados = []
    for p in doc.paragraphs:
        texto_p = limpiar_str_word(p.text).strip()
        if texto_p.startswith("tabla") and len(texto_p) < 200:
            cumple = True
            for pal in palabras_clave:
                if isinstance(pal, (list, tuple, set)):
                    if not any(limpiar_str_word(op) in texto_p for op in pal):
                        cumple = False
                        break
                else:
                    if limpiar_str_word(pal) not in texto_p:
                        cumple = False
                        break
            if cumple:
                parrafos_encontrados.append(p)
    return parrafos_encontrados

def _borrar_tabla_o_imagen_adyacente(p_elem, saltar_primera_tabla=False):
    """
    Elimina la tabla nativa o imagen de tabla vieja asociada a un párrafo 'Tabla X...'.
    """
    cursor = p_elem.getnext()
    pasos = 0
    primera_tbl_saltada = False

    while cursor is not None and pasos < 15:
        siguiente = cursor.getnext()
        tag = cursor.tag.split('}')[-1]

        if tag == 'tbl':
            if saltar_primera_tabla and not primera_tbl_saltada:
                primera_tbl_saltada = True
                cursor = siguiente
                pasos += 1
                continue
            cursor.getparent().remove(cursor)
            break

        elif tag == 'p':
            textos = [t.text for t in cursor.iter() if t.tag.endswith('}t') and t.text]
            txt_plano = "".join(textos).strip()
            # Detectar cualquier nodo de imagen en el párrafo mediante tags XML
            tiene_img = any(
                el.tag.split('}')[-1] in ('drawing', 'imagedata', 'pict', 'blip')
                for el in cursor.iter()
            )
            if tiene_img and len(txt_plano) < 10:
                cursor.getparent().remove(cursor)
                break
            elif len(txt_plano) >= 10:
                # Llegamos a un párrafo de texto normal del documento; detenemos la búsqueda
                break

        cursor = siguiente
        pasos += 1


def limpiar_tablas_viejas(parrafos_encontrados, tabla_nueva_tbl=None):
    """
    1. Borra títulos 'Tabla X...' duplicados (como en 0210 Tabla 1) junto con sus imágenes/tablas viejas.
    2. En el título principal (parrafos_encontrados[-1]), salta la tabla recién inyectada y borra
       la tabla nativa vieja o imagen de tabla que haya quedado debajo (como en 0210 Tabla 2/5 y 2502 Tabla 4).
    """
    if not parrafos_encontrados:
        return

    # 1. Títulos duplicados anteriores (no tienen tabla nueva inyectada, se borra todo)
    for p_basura in parrafos_encontrados[:-1]:
        _borrar_tabla_o_imagen_adyacente(p_basura._element, saltar_primera_tabla=False)
        xml_str = p_basura._element.xml
        if 'w:type="page"' in xml_str or 'pageBreakBefore' in xml_str:
            for t in p_basura._element.iter():
                if t.tag.endswith('}t'):
                    t.text = ''
        else:
            if p_basura._element.getparent() is not None:
                p_basura._element.getparent().remove(p_basura._element)

    # 2. Título principal donde se acaba de hacer parrafos[-1]._p.addnext(table._tbl)
    p_ancla = parrafos_encontrados[-1]._element
    _borrar_tabla_o_imagen_adyacente(p_ancla, saltar_primera_tabla=True)

def tiene_datos_reales_parquet(df_filtrado, columnas_numericas):
    """
    Verifica que el DataFrame filtrado para el acuífero realmente tenga valores distintos de 0, NaN o '-'.
    Evita que se reemplacen tablas originales válidas por tablas llenas de 0.0 o '-' (como en 0210 ETR y Salidas).
    """
    if df_filtrado is None or df_filtrado.empty:
        return False
    for col in columnas_numericas:
        if col in df_filtrado.columns:
            serie = pd.to_numeric(
                df_filtrado[col].astype(str).str.replace(',', '').str.strip().replace(['-', 'nan', 'None', ''], '0'),
                errors='coerce'
            ).fillna(0.0)
            if (serie.abs() > 1e-6).any():
                return True
    return False

def estilizar_tabla_nativa_existente(table, es_ecuacion=False):
    """
    Aplica el estilo institucional CONAGUA (#C0D7EE / #244062 / Noto Sans) a una tabla nativa
    que proviene del documento original, repitiendo su encabezado si cambia de página.
    """
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    forzar_bordes_tabla(table)
    
    num_filas = len(table.rows)
    filas_encabezado = 0

    for r_idx, row in enumerate(table.rows):
        es_encabezado = (r_idx == 0) and (num_filas > 1) and not es_ecuacion
        if r_idx == 1 and num_filas > 2 and not es_ecuacion:
            if any('gridSpan' in c._tc.xml or 'vMerge' in c._tc.xml for c in table.rows[0].cells):
                es_encabezado = True

        if es_encabezado:
            filas_encabezado = r_idx + 1

        celdas_unicas = []
        vistas = set()
        for c in row.cells:
            if c._tc not in vistas:
                vistas.add(c._tc)
                celdas_unicas.append(c)

        texto_primera_celda = celdas_unicas[0].text.strip().upper() if celdas_unicas else ""
        es_fila_total = any(k in texto_primera_celda for k in [
            "TOTAL", "PROMEDIO", "EVOLUCIÓN EN EL PERÍODO", "EVOLUCION EN EL PERIODO",
            "EVOLUCIÓN MEDIA", "EVOLUCION MEDIA", "EVAPOTRANSPIRACIÓN TOTAL",
            "EVAPOTRANSPIRACION TOTAL", "ÁREA DE BALANCE", "AREA DE BALANCE"
        ]) and (r_idx >= num_filas - 2) and (num_filas > 2)

        for c_idx, cell in enumerate(celdas_unicas):
            txt = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            if es_encabezado:
                dar_formato_celda(cell, txt, fondo_color="C0D7EE", negrita=True, tamano=9, color_texto="244062")
            elif es_fila_total:
                alineacion_tot = WD_ALIGN_PARAGRAPH.RIGHT if (c_idx == 0 and len(celdas_unicas) == 2) else WD_ALIGN_PARAGRAPH.CENTER
                dar_formato_celda(cell, txt, fondo_color="C0D7EE", negrita=True, tamano=9, alineacion=alineacion_tot, color_texto="244062")
            else:
                dar_formato_celda(cell, txt, fondo_color=None, negrita=es_ecuacion, tamano=9)

    if not es_ecuacion and filas_encabezado > 0:
        configurar_encabezado_repetible_y_paginacion(table, num_filas_encabezado=filas_encabezado)

def activar_ortografia_y_gramatica_es_mx(doc):
    """
    1. Fuerza el idioma Español (México) 'es-MX' en estilos, párrafos y tablas.
    2. Elimina cualquier etiqueta <w:noProof/> que bloquee el corrector de Word.
    3. Obliga a Word (en settings.xml) a ejecutar la revisión ortográfica y gramatical al abrir el archivo.
    """
    # 1. Activar revisión en settings.xml y quitar ocultamiento de errores
    settings = doc.settings.element
    for tag_ocultar in ['hideSpellingErrors', 'hideGrammaticalErrors']:
        el = settings.find(qn(f'w:{tag_ocultar}'))
        if el is not None:
            settings.remove(el)

    proof_state = settings.find(qn('w:proofState'))
    if proof_state is None:
        proof_state = OxmlElement('w:proofState')
        settings.append(proof_state)
    # 'dirty' obliga a Word a re-analizar toda la ortografía y gramática en cuanto se abre el .docx
    proof_state.set(qn('w:spelling'), 'dirty')
    proof_state.set(qn('w:grammar'), 'dirty')

    # 2. Configurar idioma por defecto del documento en estilos (Normal y docDefaults)
    styles_elem = doc.styles.element
    for rPr_default in styles_elem.xpath('.//w:docDefaults//w:rPrDefault//w:rPr | .//w:style//w:rPr'):
        no_proof = rPr_default.find(qn('w:noProof'))
        if no_proof is not None:
            rPr_default.remove(no_proof)
        lang = rPr_default.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang')
            rPr_default.append(lang)
        lang.set(qn('w:val'), 'es-MX')
        lang.set(qn('w:eastAsia'), 'es-MX')
        lang.set(qn('w:bidi'), 'ar-SA')

    # 3. Limpiar <w:noProof/> y fijar 'es-MX' en todos los runs del cuerpo y tablas
    for r_elem in doc.element.body.xpath('.//w:r'):
        rPr = r_elem.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            r_elem.insert(0, rPr)
        no_proof = rPr.find(qn('w:noProof'))
        if no_proof is not None:
            rPr.remove(no_proof)
        lang = rPr.find(qn('w:lang'))
        if lang is None:
            lang = OxmlElement('w:lang')
            rPr.append(lang)
        lang.set(qn('w:val'), 'es-MX')
        
